"""Genera la Carta Gantt del proyecto en Excel (informe/Carta_Gantt.xlsx).

Las barras se pintan con formato condicional a partir de las columnas Inicio y Fin:
si se cambia una fecha, la barra se mueve sola. Para pegarla en Word: seleccionar A1:U15,
copiar y pegar como imagen (o como tabla).
"""
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

SALIDA = Path(__file__).resolve().parent.parent / "Carta_Gantt.xlsx"

INICIO_LINEA = date(2026, 10, 5)   # primer día de la escala de tiempo
N_DIAS = 15                         # 05-10 al 19-10

# (etapa, actividad, responsable, inicio, fin, estado)
TAREAS = [
    ("Planificación", "Lectura de la guía y definición de la Opción C", "Ambos", date(2026, 10, 5), date(2026, 10, 6), "Completada"),
    ("Implementación", "Pipeline original (notebooks 01 a 03)", "Luis", date(2026, 10, 14), date(2026, 10, 16), "Completada"),
    ("Verificación", "Reproducción, comparación y reescritura de notebooks", "Richard", date(2026, 10, 16), date(2026, 10, 17), "Completada"),
    ("", "Depuración y validación cruzada (notebook 04)", "Richard", date(2026, 10, 17), date(2026, 10, 17), "Completada"),
    ("Informe y entrega", "Diseño, planificación y capítulos propios del informe", "Richard", date(2026, 10, 17), date(2026, 10, 19), "En curso"),
    ("", "Evidencias y capítulos de implementación del informe", "Luis", date(2026, 10, 17), date(2026, 10, 19), "En curso"),
    ("", "Repositorio de GitHub", "Ambos", date(2026, 10, 18), date(2026, 10, 19), "Pendiente"),
    ("", "Entrega en el ambiente de aprendizaje (antes de las 18:00)", "Ambos", date(2026, 10, 19), date(2026, 10, 19), "Pendiente"),
]
GRUPOS = [(6, 6), (7, 7), (8, 9), (10, 13)]   # filas combinadas de la columna Etapa

COLOR = {"Richard": "2F6DB5", "Luis": "E07B00", "Ambos": "7A4FBF", "Hito": "C0392B"}
OSCURO = "232F3E"
FIN_SEMANA = "EDEDED"
fino = Side(style="thin", color="D0D0D0")
BORDE = Border(left=fino, right=fino, top=fino, bottom=fino)

wb = Workbook()
ws = wb.active
ws.title = "Carta Gantt"
ws.sheet_view.showGridLines = False

# --- Título ---------------------------------------------------------------
ws.merge_cells("A1:U1")
ws["A1"] = "Carta Gantt del proyecto: Evaluación Sumativa 02, Big Data (Opción C)"
ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=OSCURO)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[1].height = 26

# --- Encabezados ----------------------------------------------------------
PRIMERA_COL_DIA = 7  # G
cabeceras = ["Etapa", "Actividad", "Responsable", "Inicio", "Fin", "Estado"]
for i, texto in enumerate(cabeceras, start=1):
    ws.merge_cells(start_row=3, start_column=i, end_row=5, end_column=i)
    c = ws.cell(row=3, column=i, value=texto)
    c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=OSCURO)
    c.alignment = Alignment(horizontal="center", vertical="center")
    for r in (3, 4, 5):
        ws.cell(row=r, column=i).fill = PatternFill("solid", fgColor=OSCURO)

ultima_col_dia = PRIMERA_COL_DIA + N_DIAS - 1
ws.merge_cells(start_row=3, start_column=PRIMERA_COL_DIA, end_row=3, end_column=ultima_col_dia)
c = ws.cell(row=3, column=PRIMERA_COL_DIA, value="Octubre de 2026")
c.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
c.alignment = Alignment(horizontal="center", vertical="center")
for col in range(PRIMERA_COL_DIA, ultima_col_dia + 1):
    L = get_column_letter(col)
    ws.cell(row=3, column=col).fill = PatternFill("solid", fgColor=OSCURO)
    # fila 5: fecha real (la primera es un valor, las demás suman un día); fila 4: día de la semana
    ws.cell(row=5, column=col, value=INICIO_LINEA if col == PRIMERA_COL_DIA else f"={get_column_letter(col - 1)}5+1")
    ws.cell(row=4, column=col, value=f"={L}5")
    ws.cell(row=5, column=col).number_format = "d"
    ws.cell(row=4, column=col).number_format = "ddd"
    for r, tam in ((4, 8), (5, 10)):
        x = ws.cell(row=r, column=col)
        x.font = Font(name="Calibri", size=tam, bold=(r == 5), color="FFFFFF")
        x.fill = PatternFill("solid", fgColor=OSCURO)
        x.alignment = Alignment(horizontal="center", vertical="center")
    ws.column_dimensions[L].width = 4.6

# --- Filas de tareas ------------------------------------------------------
fila0 = 6
for i, (etapa, actividad, resp, ini, fin, estado) in enumerate(TAREAS):
    r = fila0 + i
    valores = [etapa, actividad, resp, ini, fin, estado]
    for col, v in enumerate(valores, start=1):
        c = ws.cell(row=r, column=col, value=v if v != "" else None)
        c.font = Font(name="Calibri", size=10, bold=(col == 1))
        c.border = BORDE
        c.alignment = Alignment(horizontal="center" if col >= 3 else "left", vertical="center", wrap_text=(col <= 2))
        if col in (4, 5):
            c.number_format = "dd-mm"
    for col in range(PRIMERA_COL_DIA, ultima_col_dia + 1):
        ws.cell(row=r, column=col).border = BORDE
    ws.row_dimensions[r].height = 24
for a, b in GRUPOS:
    if b > a:
        ws.merge_cells(start_row=a, start_column=1, end_row=b, end_column=1)
    ws.cell(row=a, column=1).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

ultima_fila = fila0 + len(TAREAS) - 1
ws.column_dimensions["A"].width = 17
ws.column_dimensions["B"].width = 50
ws.column_dimensions["C"].width = 12
ws.column_dimensions["D"].width = 8
ws.column_dimensions["E"].width = 8
ws.column_dimensions["F"].width = 12

# --- Formato condicional --------------------------------------------------
rango = f"{get_column_letter(PRIMERA_COL_DIA)}{fila0}:{get_column_letter(ultima_col_dia)}{ultima_fila}"
G = get_column_letter(PRIMERA_COL_DIA)
dentro = f"AND({G}$5>=$D{fila0},{G}$5<=$E{fila0})"
# La última fila es el hito de entrega (rojo); el resto se colorea según el responsable.
rango_hito = f"{G}{ultima_fila}:{get_column_letter(ultima_col_dia)}{ultima_fila}"
ws.conditional_formatting.add(rango_hito, FormulaRule(
    formula=[f"AND({G}$5>=$D{ultima_fila},{G}$5<=$E{ultima_fila})"],
    fill=PatternFill("solid", bgColor=COLOR["Hito"], fgColor=COLOR["Hito"]), stopIfTrue=True))
for nombre in ("Richard", "Luis", "Ambos"):
    ws.conditional_formatting.add(rango, FormulaRule(
        formula=[f'AND({dentro},$C{fila0}="{nombre}")'],
        fill=PatternFill("solid", bgColor=COLOR[nombre], fgColor=COLOR[nombre]), stopIfTrue=True))
ws.conditional_formatting.add(rango, FormulaRule(
    formula=[f"WEEKDAY({G}$5,2)>5"],
    fill=PatternFill("solid", bgColor=FIN_SEMANA, fgColor=FIN_SEMANA)))
# Estado
rango_estado = f"F{fila0}:F{ultima_fila}"
for texto, color in (("Completada", "D5E8D4"), ("En curso", "FFF2CC"), ("Pendiente", "F8D7DA")):
    ws.conditional_formatting.add(rango_estado, FormulaRule(
        formula=[f'$F{fila0}="{texto}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))

# --- Leyenda --------------------------------------------------------------
fila_ley = ultima_fila + 2
ws.merge_cells(start_row=fila_ley, start_column=1, end_row=fila_ley, end_column=ultima_col_dia)


def cuadro(color):
    return TextBlock(InlineFont(rFont="Calibri", sz=12, color=color), "■ ")


def texto(t):
    return TextBlock(InlineFont(rFont="Calibri", sz=10), t)


ws.cell(row=fila_ley, column=1).value = CellRichText(
    texto("Leyenda:   "),
    cuadro(COLOR["Richard"]), texto("Richard      "),
    cuadro(COLOR["Luis"]), texto("Luis      "),
    cuadro(COLOR["Ambos"]), texto("Ambos      "),
    cuadro(COLOR["Hito"]), texto("Hito de entrega      "),
    cuadro(FIN_SEMANA), texto("Fin de semana"),
)
ws.cell(row=fila_ley, column=1).alignment = Alignment(horizontal="left", vertical="center")
ws.row_dimensions[fila_ley].height = 22

# --- Impresión / pegado en Word --------------------------------------------
ws.freeze_panes = f"{G}{fila0}"
ws.print_area = f"A1:{get_column_letter(ultima_col_dia)}{fila_ley}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_options.horizontalCentered = True
ws.page_margins.left = ws.page_margins.right = 0.4

wb.save(SALIDA)
print("Guardado:", SALIDA)
